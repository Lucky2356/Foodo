"""Rosstat «Средние потребительские цены» → regional and city price observations.

Source: https://rosstat.gov.ru/statistics/price, files sred_potreb_cen_<MM>-<YYYY>.xlsx (current
year) and sred_potreb_cen_<YYYY>.xlsx (past years). One sheet per month «MM(YYYY)»: territory code
and name in the first two columns, goods codes in row 4 and names with units in row 5, prices
«на конец периода, рублей за единицу измерения». These are averages, not store prices: the app
marks them «≈ оценка» and they never override a chain price or the user's own price."""

from __future__ import annotations

import io
import re
from calendar import monthrange
from typing import Any

from .common import Fetcher

PAGE = "https://rosstat.gov.ru/statistics/price"

# Rosstat goods code → (Foodo products, package size, package unit). Only goods that correspond to
# a Foodo product are used; a code may serve several products of the same kind (e.g. any fish
# fillet). Products without a fair Rosstat equivalent (chicken breast, parmesan, spices…) are left
# without an estimate instead of borrowing a wrong one.
GOODS: dict[str, tuple[tuple[str, ...], float, str]] = {
    "112": (("beef",), 1, "kg"),
    "117": (("pork", "pork_loin"), 1, "kg"),
    "116": (("lamb",), 1, "kg"),
    "115": (("chicken_thighs", "chicken_drumsticks"), 1, "kg"),
    "119": (("turkey_mince",), 1, "kg"),
    "105": (("beef_mince", "mince_pork_beef"), 1, "kg"),
    "202": (("frankfurters",), 1, "kg"),
    "204": (("smoked_sausage",), 1, "kg"),
    "208": (("bacon", "pork_ribs_smoked"), 1, "kg"),
    "404": (("cod_fillet", "pollock_fillet", "hake_fillet"), 1, "kg"),
    "403": (("salmon_fillet", "pink_salmon"), 1, "kg"),
    "411": (("mackerel",), 1, "kg"),
    "502": (("herring_salted",), 1, "kg"),
    "415": (("squid",), 1, "kg"),
    "604": (("tuna_canned", "saury_canned"), 1, "kg"),
    "701": (("butter",), 1, "kg"),
    "803": (("sunflower_oil",), 1, "l"),
    "802": (("olive_oil",), 1, "l"),
    "1111": (("milk",), 1, "l"),
    "1102": (("sour_cream",), 1, "kg"),
    "1103": (("kefir", "ryazhenka"), 1, "kg"),
    "1125": (("yogurt_plain",), 1, "kg"),
    "1116": (("cream_10", "cream_20"), 1, "l"),
    "1124": (("cottage_cheese",), 1, "kg"),
    "1204": (("cheese_hard", "mozzarella"), 1, "kg"),
    "1203": (("feta",), 1, "kg"),
    "1306": (("green_peas_canned",), 1, "kg"),
    "1304": (("canned_tomatoes",), 1, "kg"),
    "1302": (("pickled_cucumbers", "corn_canned"), 1, "kg"),
    "1501": (("eggs",), 10, "pcs"),
    "1601": (("sugar",), 1, "kg"),
    "1801": (("jam",), 1, "kg"),
    "1802": (("honey",), 1, "kg"),
    "1708": (("cocoa_powder",), 1, "kg"),
    "2001": (("mayonnaise",), 1, "kg"),
    "1305": (("ketchup",), 1, "kg"),
    "2002": (("salt",), 1, "kg"),
    "2101": (("wheat_flour",), 1, "kg"),
    "2201": (("rye_bread",), 1, "kg"),
    "2207": (("bread_white",), 1, "kg"),
    "2301": (("rice_white_dry", "rice_round_dry"), 1, "kg"),
    "2302": (("semolina",), 1, "kg"),
    "2303": (("millet_dry",), 1, "kg"),
    "2306": (("split_peas_dry", "beans_dry"), 1, "kg"),
    "2307": (("buckwheat_dry",), 1, "kg"),
    "2308": (("pearl_barley_dry",), 1, "kg"),
    "2309": (("oats_flakes",), 1, "kg"),
    "2401": (("vermicelli_dry",), 1, "kg"),
    "2403": (("pasta_dry", "spaghetti_dry"), 1, "kg"),
    "2501": (("potato",), 1, "kg"),
    "2601": (("cabbage",), 1, "kg"),
    "2603": (("onion",), 1, "kg"),
    "2604": (("beet",), 1, "kg"),
    "2605": (("carrot",), 1, "kg"),
    "2611": (("garlic",), 1, "kg"),
    "2621": (("cucumber",), 1, "kg"),
    "2623": (("tomato",), 1, "kg"),
    "2632": (("bell_pepper",), 1, "kg"),
    "2602": (("dill", "parsley", "cilantro", "green_onion"), 1, "kg"),
    "2624": (("champignons",), 1, "kg"),
    "2631": (("frozen_veg_mix", "green_beans_frozen", "green_peas_frozen"), 1, "kg"),
    "2701": (("apple",), 1, "kg"),
    "2702": (("pear",), 1, "kg"),
    "2703": (("orange",), 1, "kg"),
    "2705": (("lemon",), 1, "kg"),
    "2711": (("banana",), 1, "kg"),
    "2706": (("raisins", "prunes", "dried_apricots"), 1, "kg"),
    "2712": (("walnuts",), 1, "kg"),
    "2713": (
        ("blueberries_frozen", "strawberries_frozen", "cranberries_frozen"),
        1,
        "kg",
    ),
}

# Loose goods are costed by weight; packaged ones (eggs, milk, oil) by their unit.
LOOSE_UNITS = {"kg"}

# Rosstat territory names that differ from the region names in shared/data/geo/cities.json.
REGION_ALIASES = {
    "город москва столица российской федерации город федерального значения": "Москва",
    "город санкт-петербург город федерального значения": "Санкт-Петербург",
    "город федерального значения севастополь": "Севастополь",
    "чувашская республика - чувашия": "Чувашская Республика",
    "архангельская область (кроме ненецкого автономного округа)": "Архангельская область",
    "тюменская область (кроме ханты-мансийского автономного округа-югры и ямало-ненецкого автономного округа)": "Тюменская область",
}


def norm(name: str) -> str:
    text = str(name).lower().replace("ё", "е").replace("—", "-")
    text = re.sub(r"\s*-\s*", "-", text)
    return re.sub(r"\s+", " ", text).strip()


def region_key(name: str) -> str:
    """Normalised region name without a trailing «(…)» qualifier, e.g. «Республика Адыгея (Адыгея)»."""
    return re.sub(r"\s*\([^)]*\)$", "", norm(name))


def latest_file_url(page_html: str) -> str:
    """The newest «sred_potreb_cen» file: current-year «MM-YYYY» beats past-year «YYYY»."""
    best: tuple[int, int, str] | None = None
    for href in set(
        re.findall(r'href="([^"]*sred_potreb_cen_[^"]+\.xlsx)"', page_html)
    ):
        m = re.search(r"sred_potreb_cen_(?:(\d{2})-)?(\d{4})\.xlsx", href)
        if not m:
            continue
        key = (int(m.group(2)), int(m.group(1) or 0), href)
        if best is None or key > best:
            best = key
    if best is None:
        raise ValueError(
            "На странице Росстата не найден файл средних потребительских цен"
        )
    href = best[2]
    return href if href.startswith("http") else "https://rosstat.gov.ru" + href


def latest_month_sheet(sheet_names: list[str]) -> tuple[str, int, int]:
    months = []
    for name in sheet_names:
        m = re.fullmatch(r"(\d{2})\((\d{4})\)", name.strip())
        if m:
            months.append((int(m.group(2)), int(m.group(1)), name))
    if not months:
        raise ValueError("В файле нет листов вида «MM(YYYY)»")
    year, month, name = max(months)
    return name, year, month


def _price(value: Any) -> float | None:
    if isinstance(value, (int, float)):
        return float(value) if value > 0 else None
    try:
        number = float(str(value).replace(",", ".").replace(" ", ""))
    except ValueError:
        return None  # «...», «-» and empty cells mean «no data», never zero
    return number if number > 0 else None


def parse_sheet(
    rows: list[tuple[Any, ...]], cities: list[dict[str, Any]], year: int, month: int
) -> tuple[list[dict[str, Any]], dict[str, Any]]:
    """Observations for every mapped good, for each region and each city Rosstat reports separately."""
    header = next(
        i for i, r in enumerate(rows) if r and str(r[0]).strip() == "Код территории"
    )
    codes = [str(c).strip() if c is not None else "" for c in rows[header]]
    columns = {i: GOODS[c] for i, c in enumerate(codes) if c in GOODS}
    regions = {region_key(c["region"]): c["region"] for c in cities}
    cities_by_region: dict[str, dict[str, str]] = {}
    for c in cities:
        cities_by_region.setdefault(c["region"], {})[norm(c["name"])] = c["id"]
    observed_on = f"{year:04d}-{month:02d}-{monthrange(year, month)[1]:02d}"
    out: dict[tuple[str, str | None, str], dict[str, Any]] = {}
    stats = {"regions": set(), "cities": set(), "unmatched": []}
    region: str | None = None
    for row in rows[header + 2 :]:
        if not row or row[0] is None or row[1] is None:
            continue
        code, name = str(row[0]).strip(), str(row[1]).strip()
        if not code.isdigit():
            continue
        alias = REGION_ALIASES.get(norm(name)) or regions.get(region_key(name))
        if alias:
            region, city_id = alias, None
            stats["regions"].add(region)
        elif len(code) <= 3 or "(" in name:
            # Russia, federal districts, autonomous okrugs inside a region: not a Foodo territory.
            region = regions.get(region_key(name)) if "(" in name else None
            continue
        elif region is None:
            continue
        else:
            city_id = cities_by_region.get(region, {}).get(norm(name))
            if city_id is None:
                stats["unmatched"].append(f"{region}: {name}")
                continue
            stats["cities"].add(city_id)
        for col, (products, size, unit) in columns.items():
            value = _price(row[col]) if col < len(row) else None
            if value is None:
                continue
            for product in products:
                out[(product, city_id, region)] = {
                    "product_id": product,
                    "region": region,
                    **({"city_id": city_id} if city_id else {}),
                    "price_kop": round(value * 100),
                    "package_size": size,
                    "package_unit": unit,
                    "loose": unit in LOOSE_UNITS,
                    "observed_on": observed_on,
                    "source": f"Росстат, средние потребительские цены, {month:02d}.{year}",
                    "source_url": PAGE,
                    "external_id": f"rosstat:{code}:{codes[col]}",
                }
    summary = {
        "regions": len(stats["regions"]),
        "cities": len(stats["cities"]),
        "unmatched_cities": stats["unmatched"][:50],
        "period": f"{year:04d}-{month:02d}",
    }
    return [
        out[k] for k in sorted(out, key=lambda k: (k[2], k[1] or "", k[0]))
    ], summary


def collect(
    cities: list[dict[str, Any]], context: Any
) -> tuple[list[dict[str, Any]], dict[str, Any]]:
    import openpyxl  # only needed in the workflow

    fetcher = Fetcher("https://rosstat.gov.ru", delay=1.0, context=context, timeout=180)
    url = latest_file_url(fetcher.text(PAGE))
    book = openpyxl.load_workbook(
        io.BytesIO(fetcher.get(url)), read_only=True, data_only=True
    )
    sheet, year, month = latest_month_sheet(book.sheetnames)
    rows = [tuple(r) for r in book[sheet].iter_rows(values_only=True)]
    observations, summary = parse_sheet(rows, cities, year, month)
    summary["file"] = url
    return observations, summary

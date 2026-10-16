"""Reachability probe for price sources, run from GitHub Actions runners (round 6: Rosstat files)."""

from __future__ import annotations

import io
import re
import ssl
import sys
import urllib.error
import urllib.request

UA = "FoodoPriceBot/0.1 (+https://github.com/Lucky2356/Foodo; prices for a free meal planner)"


def fetch(url: str, ctx: ssl.SSLContext | None = None, binary: bool = False):
    req = urllib.request.Request(url, headers={"User-Agent": UA, "Accept-Language": "ru"})
    try:
        with urllib.request.urlopen(req, timeout=60, context=ctx) as r:
            data = r.read(40_000_000)
            return r.status, data if binary else data.decode("utf-8", "replace")
    except urllib.error.HTTPError as e:
        return e.code, b"" if binary else ""
    except Exception as e:  # noqa: BLE001
        return -1, f"{type(e).__name__}: {e}"


def main() -> int:
    ctx = ssl.create_default_context()
    for url in ["https://gu-st.ru/content/lending/russian_trusted_root_ca_pem.crt", "http://nuc-cdp.digital.gov.ru/cdp/subca_ssl_rsa2024.crt"]:
        _, pem = fetch(url)
        ctx.load_verify_locations(cadata=pem)
    _, page = fetch("https://rosstat.gov.ru/statistics/price", ctx)
    items = re.findall(r'document-list__item-title">\s*([^<]+?)\s*<.*?href="([^"]+\.(?:xlsx?|zip|csv|docx?))"', page, re.S)
    print(f"{len(items)} documents")
    for title, href in items:
        if re.search(r"потреб|продовольств|продукт", title, re.I):
            print(" *", title.strip()[:160], "->", href)
    import openpyxl  # installed by the workflow step

    picks = [href for title, href in items if re.search(r"средн\w* (потребительск|цен)", title, re.I) and re.search(r"потреб", title, re.I)]
    for href in picks[:3]:
        url = href if href.startswith("http") else "https://rosstat.gov.ru" + href
        status, data = fetch(url, ctx, binary=True)
        print(f"\nFILE {url} -> {status}, {len(data)} bytes")
        if status != 200 or not url.endswith("xlsx"):
            continue
        wb = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        for ws in wb.worksheets[:4]:
            print(f"  SHEET {ws.title!r} dims={ws.max_row}x{ws.max_column}")
            for i, row in enumerate(ws.iter_rows(values_only=True)):
                if i >= 12:
                    break
                print("   ", [str(c)[:40] if c is not None else "" for c in row[:10]])
    return 0


if __name__ == "__main__":
    sys.exit(main())
